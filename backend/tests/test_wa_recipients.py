import io
import zipfile
from datetime import date, datetime

import openpyxl

from app.services.wa_recipients import (
    build,
    flatten_param,
    parse_paste,
    parse_upload,
    phone10,
)


def test_phone10_survives_excel_floats_and_prefixes():
    # Review Focus #1
    assert phone10("+91 98765 43210") == "9876543210"
    assert phone10("919876543210") == "9876543210"
    assert phone10("9876543210.0") == "9876543210"
    assert phone10(9.1987654321e11) == "9876543210"
    assert phone10("98765") is None
    assert phone10("") is None


def test_phone10_never_guesses_a_number_excel_rounded():
    # Review Focus #1 — a CSV saved from Excel holds what the cell DISPLAYED
    assert phone10("9.19876543210E+11") == "9876543210"  # every digit present → exact
    assert phone10("9.19876543210E11") == "9876543210"
    assert phone10("9.19877E+11") is None                 # Excel rounded it away
    assert phone10("9.1987654321E+11") is None            # 11 of 12 digits: the last one would be a guess


def test_values_are_flattened_to_one_line_for_whatsapp():
    """Meta refuses a template param with a newline, a tab or 5+ spaces in a row (error 132018)."""
    out = build([["9876543210", "Rahul\nKumar", "Sector\t62,     Noida"]], phone_col=0, name_col=None,
                var_cols=[1, 2], fixed=[None, None], defaults=[None, None])
    assert out[0].variables == ["Rahul Kumar", "Sector 62, Noida"] and out[0].error is None


def test_csv_upload_has_a_header_row():
    headers, rows = parse_upload("list.csv", b"Phone,Name,City\n9876543210,Rahul,Noida\n")
    assert headers == ["Phone", "Name", "City"]
    assert rows == [["9876543210", "Rahul", "Noida"]]


def test_xlsx_upload_reads_the_first_sheet():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Phone", "Name"])
    ws.append([919876543210, "Rahul"])
    buf = io.BytesIO()
    wb.save(buf)
    headers, rows = parse_upload("list.xlsx", buf.getvalue())
    assert headers == ["Phone", "Name"]
    assert phone10(rows[0][0]) == "9876543210"


def test_paste_accepts_commas_or_tabs():
    assert parse_paste("9876543210, Rahul, Noida\n9876543211\tPriya\n\n") == [
        ["9876543210", "Rahul", "Noida"], ["9876543211", "Priya"]]


def test_build_maps_columns_fixed_values_and_defaults():
    rows = [["9876543210", "Rahul", ""], ["9876543211", "", "Gurgaon"]]
    out = build(rows, phone_col=0, name_col=1, var_cols=[1, 2], fixed=[None, None], defaults=["there", None])
    assert out[0].variables == ["Rahul", None]
    assert out[0].error == "missing_variable", "slot 2 blank with no default"
    assert out[1].variables == ["there", "Gurgaon"] and out[1].error is None


def test_a_number_twice_keeps_the_first_and_flags_the_second():
    # Review Focus #3
    rows = [["9876543210", "Rahul"], ["+91 98765 43210", "Rahul again"]]
    out = build(rows, phone_col=0, name_col=1, var_cols=[1], fixed=[None], defaults=[None])
    assert [c.error for c in out] == [None, "duplicate"]
    assert out[0].variables == ["Rahul"]


def test_a_fixed_value_fills_a_slot_for_everyone():
    out = build([["9876543210"]], phone_col=0, name_col=None, var_cols=[None], fixed=["Noida"], defaults=[None])
    assert out[0].variables == ["Noida"] and out[0].error is None


# Fix round 1: Critical phone10 fix tests
def test_phone10_refuses_non_indian_numbers():
    """phone10 must refuse anything not exactly one Indian mobile, never truncate."""
    assert phone10("+44 7911 123456") is None  # UK number
    assert phone10("+974 5512 3456") is None   # Qatar number
    assert phone10("9876543210 / 9876543211") is None  # two numbers
    assert phone10("9876543210/11") is None    # garbage
    assert phone10("9876543210 (Flat 402)") is None  # text suffix
    assert phone10("+919876543210.0") == "9876543210"  # valid with .0


def test_phone10_rejects_unicode_digits():
    """Devanagari and other Unicode digits are rejected."""
    assert phone10("९८७६५४३२१०") is None  # Devanagari digits


def test_phone10_validates_mobile_range():
    """Indian mobiles start with 6–9, not 0–5."""
    assert phone10("1234567890") is None  # starts with 1
    assert phone10("0000000000") is None  # all zeros
    assert phone10("5123456789") is None  # starts with 5


def test_phone10_accepts_valid_prefixes():
    """Indian mobiles can be behind 00/91/0."""
    assert phone10("+91 (0) 98765 43210") == "9876543210"
    assert phone10("0091 98765 43210") == "9876543210"
    assert phone10("9198765432") == "9198765432"  # starts with 91 but is a real number


def test_tests_pin_phone10_name_and_error():
    """Mutation: tests must assert phone10, name, and error explicitly."""
    out = build(
        [["9.19877E+11", "A"], ["+91 98765 43210", "Rahul"]],
        phone_col=0,
        name_col=1,
        var_cols=[],
        fixed=[],
        defaults=[],
    )
    assert [(c.phone10, c.name, c.error) for c in out] == [
        (None, "A", "invalid_phone"),
        ("9876543210", "Rahul", None),
    ]


def test_parse_upload_unknown_file_types_rejected():
    """Only .csv / .txt (CSV) and .xlsx / .xlsm (Excel) are accepted."""
    try:
        parse_upload("list.pdf", b"nonsense")
        assert False, "should raise ValueError"
    except ValueError as e:
        assert "upload a .csv or .xlsx file" in str(e)
    try:
        parse_upload("list.xls", b"nonsense")
        assert False, "should raise ValueError"
    except ValueError:
        pass


def test_parse_upload_csv_fallback_to_cp1252():
    """Excel's CSV export is cp1252; utf-8-sig fallback on decode error."""
    # 0x92 in cp1252 is the right single quotation mark (U+2019)
    data = b"Name\nSneha\x92s\n"
    headers, rows = parse_upload("list.csv", data)
    # cp1252 decoded: 0x92 → U+2019 (RIGHT SINGLE QUOTATION MARK)
    assert rows[0][0] == "Sneha’s"


def test_parse_upload_xlsx_date_at_midnight():
    """XLSX datetime at midnight becomes ISO date."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Date"])
    ws.append([datetime(2026, 10, 1, 0, 0, 0)])
    buf = io.BytesIO()
    wb.save(buf)
    headers, rows = parse_upload("list.xlsx", buf.getvalue())
    assert rows[0][0] == "2026-10-01"


def test_parse_upload_xlsx_date_type():
    """XLSX date (not datetime) becomes ISO date."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Date"])
    ws.append([date(2026, 10, 1)])
    buf = io.BytesIO()
    wb.save(buf)
    headers, rows = parse_upload("list.xlsx", buf.getvalue())
    assert rows[0][0] == "2026-10-01"


def test_duplicate_only_valid_rows_claim_the_number():
    """Only error-free rows add to seen; duplicates of invalid rows don't claim it."""
    rows = [["9876543210", ""], ["9876543210", "Rahul"]]
    out = build(rows, phone_col=0, name_col=1, var_cols=[1], fixed=[None], defaults=[None])
    assert [c.error for c in out] == ["missing_variable", None]
    assert out[1].variables == ["Rahul"]


def test_flatten_param_handles_none():
    """flatten_param must accept None and whitespace-only strings."""
    assert flatten_param(None) is None
    assert flatten_param("  ") is None
    assert flatten_param("Rahul") == "Rahul"
    assert flatten_param("Rahul\nKumar") == "Rahul Kumar"
    assert flatten_param("Sector\t62,     Noida") == "Sector 62, Noida"


def test_parse_upload_xlsx_stale_dimension():
    """read_only=True trusts <dimension>; reset_dimensions() recovers all rows."""
    # Build a workbook and patch its dimension tag to be stale (A1 only)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Phone", "Name"])
    ws.append(["9876543210", "Rahul"])
    ws.append(["9876543211", "Priya"])
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()

    # Patch the dimension in the XML
    with zipfile.ZipFile(io.BytesIO(data), "r") as zin:
        sheet_xml = zin.read("xl/worksheets/sheet1.xml")
    # Replace dimension tag to be stale
    sheet_xml = sheet_xml.replace(
        b'dimension ref="A1:B3"', b'dimension ref="A1"'
    )
    patched = io.BytesIO()
    with zipfile.ZipFile(patched, "w") as zout:
        for item in zipfile.ZipFile(io.BytesIO(data)).infolist():
            if item.filename == "xl/worksheets/sheet1.xml":
                zout.writestr(item, sheet_xml)
            else:
                zout.writestr(item, zipfile.ZipFile(io.BytesIO(data)).read(item.filename))
    patched.seek(0)

    headers, rows = parse_upload("list.xlsx", patched.getvalue())
    assert len(rows) == 2
    assert rows[0][0] == "9876543210"


def test_parse_upload_bom_stripped_from_header():
    """CSV BOM (UTF-8 BOM) is stripped from the first header."""
    # UTF-8 BOM is
    headers, rows = parse_upload("list.csv", "﻿Phone,Name\n9876543210,Rahul\n".encode("utf-8"))
    assert headers == ["Phone", "Name"]


def test_parse_upload_blank_rows_dropped():
    """Blank rows in the middle are skipped (CSV)."""
    headers, rows = parse_upload("list.csv", b"Phone,Name\n9876543210,Rahul\n\n9876543211,Priya\n")
    assert len(rows) == 2
    assert rows[0] == ["9876543210", "Rahul"]
    assert rows[1] == ["9876543211", "Priya"]


def test_parse_upload_blank_rows_dropped_xlsx():
    """Blank rows in the middle are skipped (xlsx)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Phone", "Name"])
    ws.append(["9876543210", "Rahul"])
    ws.append([])  # blank row in the middle
    ws.append(["9876543211", "Priya"])
    buf = io.BytesIO()
    wb.save(buf)
    headers, rows = parse_upload("list.xlsx", buf.getvalue())
    assert len(rows) == 2
    assert rows[0] == ["9876543210", "Rahul"]
    assert rows[1] == ["9876543211", "Priya"]


def test_parse_upload_xlsx_first_sheet():
    """xlsx reads the FIRST sheet, even if the active sheet is different."""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.append(["Phone", "Name"])
    ws1.append(["9876543210", "Rahul"])
    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["Other", "Data"])
    wb.active = 1  # make Sheet2 the active sheet
    buf = io.BytesIO()
    wb.save(buf)
    headers, rows = parse_upload("list.xlsx", buf.getvalue())
    assert headers == ["Phone", "Name"]
    assert rows[0] == ["9876543210", "Rahul"]


def test_parse_upload_xlsx_whole_number_float():
    """XLSX whole-number floats (2.0) render as "2", not "2.0"."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Phone", "Value"])
    ws.append([9876543210.0, 2.0])
    buf = io.BytesIO()
    wb.save(buf)
    headers, rows = parse_upload("list.xlsx", buf.getvalue())
    assert rows[0][1] == "2"


def test_cell_float_handling():
    """_cell directly: whole-number floats render "2"; fractional floats render "2.5"."""
    from app.services.wa_recipients import _cell
    assert _cell(2.0) == "2"
    assert _cell(2.5) == "2.5"


def test_cell_datetime_midnight():
    """_cell with midnight datetime renders as ISO date only."""
    from app.services.wa_recipients import _cell
    dt_midnight = datetime(2026, 10, 1, 0, 0, 0)
    assert _cell(dt_midnight) == "2026-10-01"


def test_cell_datetime_non_midnight():
    """_cell with non-midnight datetime renders as ISO with space separator."""
    from app.services.wa_recipients import _cell
    dt_non_midnight = datetime(2026, 10, 1, 9, 30, 0)
    assert _cell(dt_non_midnight) == "2026-10-01 09:30:00"


def test_cell_date():
    """_cell with date renders as ISO date."""
    from app.services.wa_recipients import _cell
    d = date(2026, 10, 1)
    assert _cell(d) == "2026-10-01"


def test_phone10_refuses_fractional_float():
    """phone10 refuses fractional floats (they round away digits)."""
    assert phone10(9876543210.7) is None
    assert phone10(9876543210.5) is None


def test_phone10_refuses_non_ascii_digits():
    """phone10 refuses non-ASCII digits (Devanagari, etc.)."""
    # Devanagari digits ०-९
    assert phone10("९८७६५४३२१०") is None
    # Mixed ASCII and Devanagari
    assert phone10("९८७६५४३२१० / 9876543211") is None
